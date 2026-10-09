# -*- coding: utf-8 -*-
"""
md_to_docx.py —— 把 docs/ 下的 Markdown 交付文档合并转换为一份 Word 文档
=====================================================================
用法：
    python tools/md_to_docx.py

输出：
    docs/全国旅游业发展可视化分析系统_项目交付文档.docx

设计说明：
  · 支持的 Markdown 语法：H1—H4 标题、段落、无序/有序列表、引用块、
    分隔线、围栏代码块（```）、管道表格、行内 **加粗** / `代码` / *斜体*；
  · 每个源文件之间自动插入分页符，便于打印与装订；
  · 中文字体统一为微软雅黑，代码块使用 Consolas 并加浅灰底纹。
"""

from __future__ import annotations

import os
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DOCS = os.path.join(ROOT, "docs")

# 合并顺序（含标题）
SOURCES = [
    ("00_交付清单与阅读指引.md", "交付清单与阅读指引"),
    ("01_项目介绍.md", "项目介绍"),
    ("02_环境安装清单.md", "环境安装清单"),
    ("03_数据集说明与样例数据.md", "数据集说明与样例数据"),
    ("04_完整源代码清单.md", "完整源代码清单"),
    ("05_运行说明.md", "运行说明"),
    ("06_效果图文字描述（论文用）.md", "效果图文字描述（论文用）"),
]

CN_FONT = "Microsoft YaHei"
CODE_FONT = "Consolas"


# ---------------------------------------------------------------------------
# 基础工具
# ---------------------------------------------------------------------------
def set_run_font(run, name=CN_FONT, size=None, bold=None, color=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), CN_FONT if name == CODE_FONT else name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def shade(element, hex_fill):
    """给段落或单元格加底纹。"""
    pr = element.get_or_add_pPr() if element.tag.endswith("}p") else element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)")


def add_inline(par, text, base_size=10.5):
    """把含行内标记的文本写入段落。"""
    for tok in INLINE_RE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            r = par.add_run(tok[2:-2]); set_run_font(r, size=base_size, bold=True)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); set_run_font(r, CODE_FONT, size=base_size - 0.5,
                                                      color=(0xC0, 0x39, 0x2B))
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); set_run_font(r, size=base_size, italic=True)
        else:
            r = par.add_run(tok); set_run_font(r, size=base_size)


# ---------------------------------------------------------------------------
# 解析与写入
# ---------------------------------------------------------------------------
def is_table_sep(line: str) -> bool:
    s = line.strip()
    if not s.startswith("|"):
        return False
    return bool(re.fullmatch(r"\|[\s:\-|]+\|", s))


def split_row(line: str):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def add_code_block(doc, lines):
    """代码块：逐行写入，等宽字体 + 浅灰底纹 + 不自动换行缩进。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(2); pf.space_after = Pt(8)
    pf.left_indent = Pt(6)
    pf.line_spacing = 1.0
    shade(p._p, "F4F6F8")
    for i, ln in enumerate(lines):
        if i:
            p.add_run().add_break()
        r = p.add_run(ln if ln else " ")
        set_run_font(r, CODE_FONT, size=8)
    return p


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        cell = t.rows[0].cells[j]
        cell.text = ""
        par = cell.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(par, h, base_size=9.5)
        for r in par.runs:
            r.font.bold = True
        shade(cell._tc, "E8F1FF")
    for row in body:
        cells = t.add_row().cells
        for j in range(len(header)):
            val = row[j] if j < len(row) else ""
            cells[j].text = ""
            add_inline(cells[j].paragraphs[0], val, base_size=9.5)
    doc.add_paragraph()
    return t


def convert_md(doc, text, add_page_break_first=False):
    lines = text.split("\n")
    i = 0
    first_block = True
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # ---- 围栏代码块 ----
        if stripped.startswith("```"):
            buf = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            add_code_block(doc, buf)
            first_block = False
            continue

        # ---- 表格 ----
        if stripped.startswith("|") and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
            rows = [split_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i])); i += 1
            add_table(doc, rows)
            first_block = False
            continue

        # ---- 标题 ----
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            title = m.group(2).strip()
            if add_page_break_first and first_block:
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            h = doc.add_heading(level=min(level, 4))
            h.paragraph_format.space_before = Pt(10 if level > 1 else 0)
            h.paragraph_format.space_after = Pt(6)
            add_inline(h, title, base_size={1: 18, 2: 15, 3: 13, 4: 11.5}.get(level, 11))
            for r in h.runs:
                r.font.bold = True
                set_run_font(r, CN_FONT,
                             size={1: 18, 2: 15, 3: 13, 4: 11.5}.get(level, 11),
                             bold=True, color=(0x0B, 0x1B, 0x33))
            first_block = False
            i += 1
            continue

        # ---- 分隔线 ----
        if re.fullmatch(r"-{3,}", stripped):
            p = doc.add_paragraph()
            pr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "BBBBBB")
            pbdr.append(bottom); pr.append(pbdr)
            p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
            first_block = False
            i += 1
            continue

        # ---- 引用块 ----
        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip()); i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            shade(p._p, "FFF8E1")
            add_inline(p, "\n".join(buf).replace("\n", " "), base_size=10)
            first_block = False
            continue

        # ---- 列表 ----
        m = re.match(r"^(\s*)([-*+]|\d+\.)\s+(.*)$", line)
        if m:
            indent = len(m.group(1)) // 2
            ordered = bool(re.match(r"\d+\.", m.group(2)))
            p = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
            p.paragraph_format.left_indent = Pt(18 + indent * 16)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(3), base_size=10.5)
            first_block = False
            i += 1
            continue

        # ---- 空行 ----
        if not stripped:
            i += 1
            continue

        # ---- 普通段落 ----
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.25
        add_inline(p, stripped, base_size=10.5)
        first_block = False
        i += 1


def add_cover(doc):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("基于 Python 的全国旅游业发展可视化分析系统")
    set_run_font(r, CN_FONT, size=26, bold=True, color=(0x0B, 0x1B, 0x33))

    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("项 目 交 付 文 档")
    set_run_font(r, CN_FONT, size=17, bold=False, color=(0x17, 0x61, 0x8F))

    for _ in range(3):
        doc.add_paragraph()

    info = [
        ("项目性质", "本科毕业设计（数据分析与可视化方向）"),
        ("技术栈", "Python 3.9+ · Pandas · NumPy · Matplotlib · Seaborn · Scikit-learn · Pyecharts 2.1.0"),
        ("部署形态", "纯静态 HTML 单文件，零后端、全离线"),
        ("数据范围", "2015—2024 年 · 全国 + 31 个省级行政区 + 20 个热门旅游城市"),
        ("文档内容", "① 项目介绍　② 环境安装清单　③ 数据集说明与样例数据　④ 完整源代码　⑤ 运行说明　⑥ 效果图文字描述"),
    ]
    for k, v in info:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(f"{k}："); set_run_font(r, CN_FONT, size=10.5, bold=True)
        r = p.add_run(v); set_run_font(r, CN_FONT, size=10.5)


def main():
    out_path = os.path.join(DOCS, "全国旅游业发展可视化分析系统_项目交付文档.docx")
    doc = Document()

    # 全局样式
    normal = doc.styles["Normal"]
    normal.font.name = CN_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    for s in doc.sections:
        s.left_margin = s.right_margin = Pt(56)
        s.top_margin = s.bottom_margin = Pt(56)

    add_cover(doc)

    missing = []
    for idx, (fname, _label) in enumerate(SOURCES):
        path = os.path.join(DOCS, fname)
        if not os.path.exists(path):
            missing.append(fname); continue
        with open(path, encoding="utf-8") as f:
            text = f.read()
        convert_md(doc, text, add_page_break_first=(idx > 0))

    doc.save(out_path)
    size_mb = os.path.getsize(out_path) / 1024 / 1024
    print(f"[完成] {out_path}")
    print(f"       大小 {size_mb:.2f} MB；段落数 {len(doc.paragraphs)}；表格数 {len(doc.tables)}")
    if missing:
        print("[缺失] " + "、".join(missing))
    return out_path


if __name__ == "__main__":
    main()
